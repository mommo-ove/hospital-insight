from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
import json
import re
from uuid import uuid4
from zipfile import ZipFile, BadZipFile

import openpyxl
from openpyxl.utils import get_column_letter
from sqlalchemy import select

from .catalog import BY_ID, HEADERS, METRICS
from .db import Database, encode, utcnow


class ImportProblem(Exception):
    def __init__(self, message, status=422):
        super().__init__(message)
        self.status = status


def parse_month(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    match = re.fullmatch(r"(\d{4})[年/\-](\d{1,2})(?:月|[-/]\d{1,2})?", str(value).strip())
    if not match:
        raise ValueError("年月必须为日期或 YYYY-MM")
    return date(int(match[1]), int(match[2]), 1).strftime("%Y-%m")


def parse_number(value, field):
    if value is None or isinstance(value, bool):
        raise ValueError("不能为空，且必须为数值（零值允许）")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        raise ValueError("必须为数值") from None
    if not number.is_finite() or number < 0 or number > Decimal("1000000000000"):
        raise ValueError("数值须为有限的非负数且不超过一万亿")
    if field == "sequence" or BY_ID[field].unit in {"张", "人次", "台次"}:
        if number != number.to_integral_value():
            raise ValueError("必须为整数")
        return int(number)
    if BY_ID[field].unit == "元":
        if number * 100 != (number * 100).to_integral_value():
            raise ValueError("金额最多支持两位小数")
        return int(number * 100)
    if BY_ID[field].unit == "%" and number > 100:
        raise ValueError("百分比应在 0 到 100 之间，如 70.9")
    return float(number)


def parse_excel(content: bytes, filename: str, import_id: str):
    if len(content) > 10 * 1024 * 1024:
        raise ImportProblem("Excel 文件不能超过 10 MB", 413)
    if not filename.lower().endswith(".xlsx"):
        raise ImportProblem("第一版只接受 .xlsx 文件")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ImportProblem("工作簿解压后超过 100 MB")
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True)
    except (BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ImportProblem("无法读取 Excel，请确认文件未损坏") from exc
    try:
        sheets = [s for s in workbook if s.max_row and s.max_row > 1]
        if len(sheets) != 1:
            raise ImportProblem("请提供仅含一张数据工作表的文件")
        sheet = sheets[0]
        if sheet.max_row > 20001 or sheet.max_column > 100:
            raise ImportProblem("最多支持 20000 条记录和 100 个工作表列")
        row_iter = sheet.iter_rows()
        header_cells = next(row_iter)
        headers = [str(c.value or "").strip().replace("（", "(").replace("）", ")") for c in header_cells]
        nonempty_headers = [h for h in headers if h]
        if len(nonempty_headers) != len(set(nonempty_headers)):
            raise ImportProblem("存在重复表头")
        missing = sorted(set(HEADERS) - set(headers))
        unknown = sorted(set(nonempty_headers) - set(HEADERS))
        if missing or unknown:
            raise ImportProblem("表头不匹配。缺少：" + "、".join(missing) + "；未知：" + "、".join(unknown))
        positions = {HEADERS[h]: i for i, h in enumerate(headers) if h in HEADERS}
        columns = {field: get_column_letter(i + 1) for field, i in positions.items()}
        rows, errors, keys, department_types = [], [], set(), {}
        digest = sha256(content).hexdigest()
        for row_number, cells in enumerate(row_iter, start=2):
            if all(c.value is None for c in cells):
                continue
            record, original, row_errors = {}, {}, []
            for field, idx in positions.items():
                cell = cells[idx]
                value = cell.value
                original[field] = value.isoformat() if isinstance(value, (date, datetime)) else value
                try:
                    if cell.data_type == "f":
                        raise ValueError("请先将公式转换为经过确认的数值")
                    if field == "month":
                        record[field] = parse_month(value)
                    elif field in {"department", "department_type"}:
                        if not isinstance(value, str) or not value.strip() or len(value.strip()) > 100:
                            raise ValueError("须为非空文字，最长100字")
                        record[field] = value.strip()
                    else:
                        record[field] = parse_number(value, field)
                except (ValueError, TypeError, OverflowError) as exc:
                    row_errors.append(f"{get_column_letter(idx+1)}{row_number}：{exc}")
            if not row_errors:
                key = (record["month"], record["department"])
                if key in keys:
                    row_errors.append(f"第 {row_number} 行：年月＋科室重复")
                keys.add(key)
                dept = record["department"]
                if dept in department_types and department_types[dept] != record["department_type"]:
                    row_errors.append(f"第 {row_number} 行：同一科室的类型不一致")
                department_types[dept] = record["department_type"]
                if sum(record[m.id] for m in METRICS[:7]) != record["total_revenue"]:
                    row_errors.append(f"第 {row_number} 行：七项收入之和与合计收入不一致")
                if record["insurance_amount"] + record["self_pay_amount"] != record["total_revenue"]:
                    row_errors.append(f"第 {row_number} 行：医保金额＋自费金额与合计收入不一致")
            errors.extend(row_errors)
            if not row_errors:
                record["source_json"] = encode({"file": filename, "sheet": sheet.title, "row": row_number,
                    "import_id": import_id, "sha256": digest, "columns": columns, "values": original})
                rows.append(record)
        if not rows and not errors:
            errors.append("工作表没有有效数据行")
        return rows, errors, sheet.title
    finally:
        workbook.close()


def same_record(left, right):
    return all(left[key] == right[key] for key in HEADERS.values())


def preview_import(db: Database, content: bytes, filename: str):
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1][:200]
    import_id = uuid4().hex
    rows, errors, sheet = parse_excel(content, filename, import_id)
    with db.write() as con:
        existing = {(r["month"], r["department"]): dict(r) for r in con.execute(select(db.facts)).mappings()}
        revision = db.revision(con)
        new, duplicate, conflict = 0, 0, 0
        conflicts = []
        for row in rows:
            before = existing.get((row["month"], row["department"]))
            if before is None:
                new += 1
            elif same_record(before, row):
                duplicate += 1
            else:
                conflict += 1
                if len(conflicts) < 50:
                    changed = []
                    for key in HEADERS.values():
                        if before[key] != row[key]:
                            money = key in BY_ID and BY_ID[key].unit == "元"
                            changed.append({"field": BY_ID[key].label if key in BY_ID else key,
                                "before": before[key] / 100 if money else before[key], "after": row[key] / 100 if money else row[key]})
                    conflicts.append({"month": row["month"], "department": row["department"], "changes": changed})
        # Existing records outside the upload must retain a consistent department classification.
        incoming_keys = {(r["month"], r["department"]) for r in rows}
        remaining_types = {r["department"]: r["department_type"] for k, r in existing.items() if k not in incoming_keys}
        for row in rows:
            if row["department"] in remaining_types and remaining_types[row["department"]] != row["department_type"]:
                errors.append(f"{row['department']} 的科室类型与历史记录不一致")
                break
        summary = {"id": import_id, "filename": filename, "sheet": sheet, "new": new, "duplicates": duplicate,
            "conflict_count": conflict, "conflicts": conflicts, "errors": errors[:100], "error_count": len(errors),
            "row_count": len(rows), "base_revision": revision, "can_commit": not errors,
            "sample": [{"month": r["month"], "department": r["department"], "total_revenue": r["total_revenue"]/100} for r in rows[:5]]}
        con.execute(db.imports.insert().values(id=import_id, filename=filename, created_at=utcnow(),
            status="invalid" if errors else "pending", base_revision=revision, payload_json=encode(rows), summary_json=encode(summary)))
    return summary


def commit_import(db: Database, import_id: str, overwrite: bool):
    with db.write() as con:
        item = con.execute(select(db.imports).where(db.imports.c.id == import_id)).mappings().first()
        if not item:
            raise ImportProblem("导入预览不存在", 404)
        if item["status"] == "committed":
            return {**json.loads(item["summary_json"]), "revision": item["revision"], "status": "committed"}
        if item["status"] != "pending":
            raise ImportProblem("校验未通过，不能提交")
        summary = json.loads(item["summary_json"])
        if summary["conflict_count"] and not overwrite:
            raise ImportProblem("存在冲突，请在预览中确认覆盖", 409)
        if db.revision(con) != item["base_revision"]:
            raise ImportProblem("预览后数据已更新，请重新上传以检查最新冲突", 409)
        rows = json.loads(item["payload_json"])
        existing = {(r["month"], r["department"]): dict(r) for r in con.execute(select(db.facts)).mappings()}
        revision = item["base_revision"] + (1 if summary["new"] or summary["conflict_count"] else 0)
        for row in rows:
            source = json.loads(row["source_json"])
            source["version"] = revision
            row["source_json"] = encode(source)
            previous = existing.get((row["month"], row["department"]))
            if previous is None:
                con.execute(db.facts.insert().values(**row))
            elif not same_record(previous, row):
                con.execute(db.facts.update().where(db.facts.c.id == previous["id"]).values(**row))
        con.execute(db.state.update().where(db.state.c.id == 1).values(revision=revision))
        con.execute(db.imports.update().where(db.imports.c.id == import_id).values(status="committed", revision=revision))
    return {**summary, "revision": revision, "status": "committed"}


def import_history(db: Database):
    with db.read_engine.connect() as con:
        rows = con.execute(select(db.imports).order_by(db.imports.c.created_at.desc()).limit(50)).mappings().all()
    return [{"id": r["id"], "filename": r["filename"], "created_at": r["created_at"], "status": r["status"],
        "revision": r["revision"], "summary": json.loads(r["summary_json"])} for r in rows]

