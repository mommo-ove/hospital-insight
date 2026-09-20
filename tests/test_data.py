from io import BytesIO
import json

import openpyxl
import pytest
from sqlalchemy import select

from backend.catalog import HEADERS, METRICS
from backend.importer import ImportProblem, commit_import, preview_import
from .conftest import ROOT


def workbook_bytes(change=None, reverse=False):
    workbook = openpyxl.load_workbook(ROOT / "院管数据.xlsx")
    sheet = workbook.active
    if change:
        change(sheet)
    if reverse:
        rows = [list(reversed(row)) for row in sheet.values]
        target = openpyxl.Workbook()
        for row in rows:
            target.active.append(row)
        workbook.close()
        workbook = target
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_import_matches_every_original_value(db):
    workbook = openpyxl.load_workbook(ROOT / "院管数据.xlsx", data_only=True)
    original = list(workbook.active.values)
    with db.read_engine.connect() as con:
        records = list(con.execute(select(db.facts).order_by(db.facts.c.sequence)).mappings())
    assert len(records) == 224
    for raw, stored in zip(original[1:], records, strict=True):
        for header, value in zip(original[0], raw, strict=True):
            field = HEADERS[header]
            if field == "month":
                expected = value.strftime("%Y-%m")
            elif field in {m.id for m in METRICS if m.unit == "元"}:
                expected = int(value * 100)
            else:
                expected = value
            assert stored[field] == expected, (stored["sequence"], field)
    workbook.close()


def test_duplicate_import_is_idempotent(db):
    preview = preview_import(db, workbook_bytes(), "repeat.xlsx")
    assert preview["duplicates"] == 224
    assert preview["new"] == preview["conflict_count"] == 0
    first = commit_import(db, preview["id"], False)
    second = commit_import(db, preview["id"], False)
    assert first == second
    assert db.catalog_state()["revision"] == 1
    assert db.catalog_state()["row_count"] == 224


def test_reordered_columns_and_source_cells(db):
    # New month avoids duplicate-source preservation and exercises actual imported coordinates.
    def change(sheet):
        for row in list(sheet.iter_rows())[1:]:
            row[1].value = row[1].value.replace(year=row[1].value.year + 3)
    preview = preview_import(db, workbook_bytes(change, reverse=True), "reordered.xlsx")
    assert preview["new"] == 224 and preview["can_commit"]
    commit_import(db, preview["id"], False)
    with db.read_engine.connect() as con:
        row = con.execute(select(db.facts).where(db.facts.c.month == "2027-01", db.facts.c.department == "心血管内科")).mappings().one()
    source = json.loads(row["source_json"])
    assert row["total_revenue"] == 209111500
    assert source["columns"]["total_revenue"] == "K"
    assert source["version"] == 2


def test_conflict_requires_explicit_confirmation_and_is_atomic(db):
    def change(sheet):
        for col in [5, 12, 14]:
            sheet.cell(2, col).value += 100
    preview = preview_import(db, workbook_bytes(change), "correction.xlsx")
    assert preview["conflict_count"] == 1
    with pytest.raises(ImportProblem, match="确认覆盖"):
        commit_import(db, preview["id"], False)
    assert db.catalog_state()["revision"] == 1
    commit_import(db, preview["id"], True)
    assert db.catalog_state()["revision"] == 2


@pytest.mark.parametrize("change,fragment", [
    (lambda s: setattr(s['E2'], 'value', None), "不能为空"),
    (lambda s: setattr(s['E2'], 'value', -1), "非负"),
    (lambda s: setattr(s['T2'], 'value', 1.2), "整数"),
    (lambda s: setattr(s['E2'], 'value', 0.001), "两位小数"),
    (lambda s: setattr(s['O2'], 'value', 120), "百分比"),
    (lambda s: setattr(s['B2'], 'value', '2025-13'), "month"),
    (lambda s: setattr(s['E2'], 'value', '=SUM(1,2)'), "公式"),
    (lambda s: setattr(s['L2'], 'value', 99), "不一致"),
])
def test_bad_input_does_not_mutate_data(db, change, fragment):
    preview = preview_import(db, workbook_bytes(change), "invalid.xlsx")
    assert not preview["can_commit"]
    assert any(fragment in e for e in preview["errors"])
    with pytest.raises(ImportProblem):
        commit_import(db, preview["id"], True)
    assert db.catalog_state()["revision"] == 1
    assert db.catalog_state()["row_count"] == 224


def test_duplicate_keys_in_one_file_rejected(db):
    preview = preview_import(db, workbook_bytes(lambda s: s.append([c.value for c in s[2]])), "duplicate.xlsx")
    assert any("年月＋科室重复" in e for e in preview["errors"])


def test_stale_preview_rejected(db):
    def change(sheet):
        for col in [5, 12, 14]:
            sheet.cell(2, col).value += 10
    first = preview_import(db, workbook_bytes(change), "first.xlsx")
    second = preview_import(db, workbook_bytes(change), "second.xlsx")
    commit_import(db, first["id"], True)
    with pytest.raises(ImportProblem, match="数据已更新"):
        commit_import(db, second["id"], True)


def test_corrupt_and_oversized_files(client):
    assert client.post('/api/imports/preview', files={'file': ('x.xlsx', b'not zip')}).status_code == 422
    assert client.post('/api/imports/preview', files={'file': ('x.xls', b'123')}).status_code == 422
    assert client.post('/api/imports/preview', files={'file': ('x.xlsx', b'x' * (10 * 1024 * 1024 + 1))}).status_code == 413

